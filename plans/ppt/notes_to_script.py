#!/usr/bin/env python3
"""notes_to_script.py · 講稿版 PPTX 的備忘稿 → `-解說腳本.md`

把每頁備忘稿（`⏱ 建議時間：…` ＋ 正文 ＋ `【講師提示】…`）重組成
`generate_narration.sh build` / `generate_voiceover.sh` 吃得下的講稿格式：

    ## slide-7 實際運行展示（1 分 30 秒）

    正文段落……

    **講師提示**：……

用法：
    .venv/bin/python plans/ppt/notes_to_script.py                 # 轉換同目錄所有 *-講稿版.pptx
    .venv/bin/python plans/ppt/notes_to_script.py <deck.pptx>      # 只轉一個
    （加 --force 才覆寫已存在的 -解說腳本.md）

備註：轉出來的是「解說腳本」，不是 TTS 專用稿；`generate_narration.sh build`
會再剝掉講師提示、產生 `-narration-script.md` 才送 TTS。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation

DECK_TITLE = "TW-SELFSELECT 台股買賣自動化系統"
TIME_RE = re.compile(r"^⏱\s*建議時間：\s*(.+)$")
TIP_RE = re.compile(r"^【講師提示】\s*")
WEEK_RE = re.compile(r"-(W\d)-")


def secs_of(label: str) -> int | None:
    """'1 分 30 秒' → 90；'2 分鐘' → 120；'30 秒' → 30"""
    label = label.replace(" ", "")
    m = re.fullmatch(r"(?:(\d+)分鐘?)?(?:(\d+)秒)?", label)
    if not m or not (m.group(1) or m.group(2)):
        return None
    return int(m.group(1) or 0) * 60 + int(m.group(2) or 0)


def fmt_total(total: int) -> str:
    minutes, secs = divmod(total, 60)
    if secs == 0:
        return f"{minutes} 分鐘"
    return f"{minutes} 分 {secs} 秒"


def slide_title(slide) -> str:
    """優先取標題版面配置區；沒有就取字級最大的文字框第一行。

    標題只給人看（TTS 不唸標題），所以清掉尾端標點並切掉問號／箭頭後的補充說明，
    最多留 28 字。
    """
    raw = ""
    try:
        if slide.shapes.title is not None:
            raw = slide.shapes.title.text
    except Exception:
        raw = ""
    if not raw.strip():
        best, best_size = "", 0.0
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            # 只取第一行：很多頁的標題與內文在同一個文字框
            first_line = shape.text_frame.text.split("\n")[0]
            best, best_size = _pick_title(shape, first_line, best, best_size)
        raw = best
    raw = " ".join(raw.split())
    raw = re.split(r"[?？]|=>|→", raw)[0]
    raw = raw.strip().strip(":：、,，;；")
    return raw[:28].strip()


def _pick_title(shape, first_line: str, best: str, best_size: float) -> tuple[str, float]:
    text = " ".join(first_line.split())
    if not text:
        return best, best_size
    size = max(
        (run.font.size.pt for para in shape.text_frame.paragraphs
         for run in para.runs if run.font.size),
        default=0.0,
    )
    top = shape.top if shape.top is not None else 10**9
    if size > best_size or (size == best_size and best == "" and top < 10**9):
        return text, size
    return best, best_size


def parse_notes(text: str) -> tuple[str, list[str], str]:
    """→ (建議時間標籤, 正文段落, 講師提示)"""
    lines = [ln.rstrip() for ln in text.split("\n")]
    time_label, body, tips = "", [], []
    started = False
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if body and body[-1] != "":
                body.append("")
            continue
        if not started:
            m = TIME_RE.match(stripped)
            if m:
                time_label = m.group(1).strip()
                started = True
                continue
        m = TIP_RE.match(stripped)
        if m:
            tips.append(stripped[m.end():].strip())
            continue
        if stripped.startswith(">") or stripped in ("---", "***", "___"):
            continue
        body.append(stripped)
        started = True
    while body and body[-1] == "":
        body.pop()
    return time_label, body, "\n".join(t for t in tips if t)


def convert(deck: Path) -> Path:
    prs = Presentation(deck)
    tag = WEEK_RE.search(deck.name)
    week = tag.group(1) if tag else ""
    out = deck.with_name(deck.name.replace("-講稿版.pptx", "-解說腳本.md"))

    sections, rows, total = [], [], 0
    for num, slide in enumerate(prs.slides, 1):
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        label, body, tips = parse_notes(notes)
        if not body and not tips:
            raise SystemExit(f"❌ 第 {num} 頁沒有備忘稿：{deck.name}")
        secs = secs_of(label) if label else None
        if secs:
            total += secs
        title = slide_title(slide)
        # 注意：build 的解析式是 `^## slide-(\d+)\s+(.+)$`，編號後一定要有空白＋標題
        head = f"## slide-{num} {title or f'第 {num} 頁'}（{label or '未標時間'}）"

        chunk = [head, ""]
        chunk.extend(body if body else ["（本頁無講稿）"])
        if tips:
            chunk += ["", f"**講師提示**：{tips.replace(chr(10), ' ')}"]
        sections.append("\n".join(chunk))
        rows.append((num, title, label or "-", fmt_total(total)))

    header = [
        f"# 《{DECK_TITLE}》{week} 解說腳本",
        "",
        f"> **講師稿 · 第 {week.lstrip('W')} 週課程**",
        f"> 建議時間合計約 **{fmt_total(total)}**（依各頁備忘稿估算）",
        "> 使用方式：逐頁照讀或按重點提示即興發揮；「講師提示」不需唸出。",
        "> 本檔由 `plans/ppt/notes_to_script.py` 從 `-講稿版.pptx` 備忘稿反推產生，"
        "可直接餵給 `plans/generate_narration.sh build` 重新配音。",
        "",
        "---",
        "",
    ]
    body_md = "\n\n---\n\n".join(sections)
    appendix = [
        "",
        "---",
        "",
        "## 附錄：時間分配",
        "",
        "| 頁碼 | 標題 | 建議時間 | 累計 |",
        "| ---: | --- | ---: | ---: |",
    ]
    appendix += [f"| {n} | {t or '—'} | {lab} | 累計 {cum} |" for n, t, lab, cum in rows]
    appendix += ["", f"> 合計約 {fmt_total(total)}（不含操作示範緩衝）。", ""]

    out.write_text("\n".join(header) + "\n" + body_md + "\n" + "\n".join(appendix), encoding="utf-8")
    print(f"✅ {out.name}：{len(sections)} 頁，合計 {fmt_total(total)}")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description="講稿版 PPTX 備忘稿 → -解說腳本.md")
    ap.add_argument("decks", nargs="*", type=Path)
    ap.add_argument("--force", action="store_true", help="覆寫已存在的 -解說腳本.md")
    args = ap.parse_args()

    decks = args.decks or sorted(Path(__file__).resolve().parent.glob("*-講稿版.pptx"))
    if not decks:
        print("❌ 找不到 *-講稿版.pptx", file=sys.stderr)
        return 1
    skipped = 0
    for deck in decks:
        if not deck.is_file():
            print(f"❌ 檔案不存在: {deck}", file=sys.stderr)
            return 1
        target = deck.with_name(deck.name.replace("-講稿版.pptx", "-解說腳本.md"))
        if target.exists() and not args.force:
            print(f"⏭ {target.name} 已存在（加 --force 覆寫）")
            skipped += 1
            continue
        convert(deck)
    if skipped:
        print(f"（跳過 {skipped} 個已存在的檔案）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
